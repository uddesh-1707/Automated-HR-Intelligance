from io import BytesIO

from docx import Document

from app.services.resume_parser import extract_text, parse_resume


def test_text_resume_extracts_only_present_details():
    content = b"Jordan Example\njordan@example.com\n6 years of experience\nPython, FastAPI"
    parsed = parse_resume(extract_text("resume.txt", content))
    assert parsed["email"] == "jordan@example.com"
    assert parsed["experience_years"] == 6
    assert "python" in parsed["skills"]
    assert parsed["education"] is None


def test_docx_resume_is_supported():
    document = Document()
    document.add_paragraph("Morgan Example")
    document.add_paragraph("morgan@example.com")
    stream = BytesIO()
    document.save(stream)
    text = extract_text("resume.docx", stream.getvalue())
    assert "Morgan Example" in text

