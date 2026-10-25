from io import BytesIO
from types import SimpleNamespace

from docx import Document

from app.api.v1 import recruitment
from app.services import recruitment_agent
from app.services.llm import LLMResponse, ToolCall
from app.services.email_notifications import EmailDeliveryError, EmailDeliveryRejected, EmailNotConfigured


def make_docx() -> bytes:
    document = Document()
    document.add_paragraph("Taylor Candidate")
    document.add_paragraph("taylor@example.com")
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def create_application(client, admin_headers, monkeypatch) -> str:
    monkeypatch.setattr(recruitment, "process_application", lambda *_args: None)
    job_response = client.post(
        "/api/v1/recruitment/jobs",
        headers=admin_headers,
        json={"title": "Communications Engineer", "description": "Build reliable software."},
    )
    assert job_response.status_code == 201, job_response.text
    publish_response = client.post(
        f"/api/v1/recruitment/jobs/{job_response.json()['id']}/publish",
        headers=admin_headers,
    )
    assert publish_response.status_code == 200, publish_response.text
    token = publish_response.json()["public_url"].rsplit("/", 1)[1]
    application_response = client.post(
        f"/api/v1/recruitment/public/jobs/{token}/applications",
        data={
            "name": "Taylor Candidate",
            "email": "taylor@example.com",
            "phone": "",
            "education": "",
            "experience_years": "",
            "answers": "{}",
        },
        files={"resume": ("resume.docx", make_docx(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert application_response.status_code == 202, application_response.text
    return application_response.json()["application_id"]


def mark_rejected(client, headers, application_id):
    application = client.get(
        f"/api/v1/recruitment/applications/{application_id}",
        headers=headers,
    ).json()
    candidate = application["candidate"]
    response = client.patch(
        f"/api/v1/recruitment/applications/{application_id}",
        headers=headers,
        json={
            "name": candidate["name"],
            "email": candidate["email"],
            "phone": candidate["phone"],
            "skills": candidate["skills"],
            "experience_years": candidate["experience_years"],
            "education": candidate["education"],
            "stage": "rejected",
            "hr_notes": "Decision recorded by HR for communications test.",
        },
    )
    assert response.status_code == 200, response.text


def test_email_draft_is_reviewable_and_send_attempt_is_idempotent(client, admin_headers, monkeypatch):
    application_id = create_application(client, admin_headers, monkeypatch)
    draft_url = f"/api/v1/recruitment/applications/{application_id}/email-drafts"
    draft_response = client.post(
        draft_url,
        headers=admin_headers,
        json={"email_type": "shortlist", "message": "Please select a time for a technical discussion."},
    )
    assert draft_response.status_code == 201, draft_response.text
    draft = draft_response.json()
    assert draft["recipient"] == "taylor@example.com"
    assert "technical discussion" in draft["body"]

    updated_response = client.post(
        draft_url,
        headers=admin_headers,
        json={"email_type": "shortlist", "message": "Please reply with your availability."},
    )
    assert updated_response.status_code == 201
    assert updated_response.json()["_id"] == draft["_id"]
    assert "reply with your availability" in updated_response.json()["body"]

    delivered = []
    monkeypatch.setattr(recruitment, "send_email", lambda **kwargs: delivered.append(kwargs))
    sent = client.post(
        f"/api/v1/recruitment/email-drafts/{draft['_id']}/approve-and-send",
        headers=admin_headers,
    )
    assert sent.status_code == 200
    assert sent.json()["delivery_status"] == "sent"
    assert len(delivered) == 1
    repeated_send = client.post(
        f"/api/v1/recruitment/email-drafts/{draft['_id']}/approve-and-send",
        headers=admin_headers,
    )
    assert repeated_send.status_code == 409
    assert len(delivered) == 1
    assert client.patch(
        f"/api/v1/recruitment/email-drafts/{draft['_id']}",
        headers=admin_headers,
        json={"subject": "Changed after send", "body": "Changed after send."},
    ).status_code == 409
    assert client.delete(
        f"/api/v1/recruitment/email-drafts/{draft['_id']}",
        headers=admin_headers,
    ).status_code == 409


def test_email_send_without_smtp_is_reported_and_not_marked_sent(client, admin_headers, monkeypatch):
    application_id = create_application(client, admin_headers, monkeypatch)
    mark_rejected(client, admin_headers, application_id)
    draft = client.post(
        f"/api/v1/recruitment/applications/{application_id}/email-drafts",
        headers=admin_headers,
        json={"email_type": "rejection", "message": ""},
    ).json()
    monkeypatch.setattr(
        recruitment,
        "send_email",
        lambda **_kwargs: (_ for _ in ()).throw(EmailNotConfigured("SMTP is not configured.")),
    )

    response = client.post(
        f"/api/v1/recruitment/email-drafts/{draft['_id']}/approve-and-send",
        headers=admin_headers,
    )
    assert response.status_code == 503
    assert response.json()["detail"] == "SMTP is not configured."
    assert client.post(
        f"/api/v1/recruitment/applications/{application_id}/email-drafts",
        headers=admin_headers,
        json={"email_type": "rejection", "message": ""},
    ).status_code == 409
    assert client.get(
        f"/api/v1/recruitment/email-drafts/{draft['_id']}/status",
        headers=admin_headers,
    ).json()["delivery_status"] == "not_configured"


def test_ambiguous_smtp_failure_is_not_reported_as_sent_or_retried(client, admin_headers, monkeypatch):
    application_id = create_application(client, admin_headers, monkeypatch)
    draft = client.post(
        f"/api/v1/recruitment/applications/{application_id}/email-drafts",
        headers=admin_headers,
        json={"email_type": "shortlist", "message": ""},
    ).json()
    monkeypatch.setattr(
        recruitment,
        "send_email",
        lambda **_kwargs: (_ for _ in ()).throw(EmailDeliveryError("Provider response was interrupted.")),
    )
    response = client.post(
        f"/api/v1/recruitment/email-drafts/{draft['_id']}/approve-and-send",
        headers=admin_headers,
    )
    assert response.status_code == 502
    assert "did not confirm delivery" in response.json()["detail"]
    status_response = client.get(
        f"/api/v1/recruitment/email-drafts/{draft['_id']}/status",
        headers=admin_headers,
    )
    assert status_response.json()["delivery_status"] == "unknown"
    assert client.post(
        f"/api/v1/recruitment/email-drafts/{draft['_id']}/approve-and-send",
        headers=admin_headers,
    ).status_code == 409
    assert client.post(
        f"/api/v1/recruitment/email-drafts/{draft['_id']}/retry",
        headers=admin_headers,
    ).status_code == 409


def test_explicit_smtp_rejection_allows_one_approved_retry(client, admin_headers, monkeypatch):
    application_id = create_application(client, admin_headers, monkeypatch)
    draft = client.post(
        f"/api/v1/recruitment/applications/{application_id}/email-drafts",
        headers=admin_headers,
        json={"email_type": "shortlist", "message": ""},
    ).json()
    monkeypatch.setattr(
        recruitment,
        "send_email",
        lambda **_kwargs: (_ for _ in ()).throw(EmailDeliveryRejected("The recipient was explicitly refused.")),
    )
    assert client.post(
        f"/api/v1/recruitment/email-drafts/{draft['_id']}/approve-and-send",
        headers=admin_headers,
    ).status_code == 502
    retry_response = client.post(
        f"/api/v1/recruitment/email-drafts/{draft['_id']}/retry",
        headers=admin_headers,
    )
    assert retry_response.status_code == 201, retry_response.text
    retry = retry_response.json()
    assert retry["retry_of"] == draft["_id"]
    assert retry["delivery_status"] == "not_sent"
    repeated = client.post(
        f"/api/v1/recruitment/email-drafts/{draft['_id']}/retry",
        headers=admin_headers,
    )
    assert repeated.status_code == 201
    assert repeated.json()["_id"] == retry["_id"]
    monkeypatch.setattr(recruitment, "send_email", lambda **_kwargs: None)
    sent = client.post(
        f"/api/v1/recruitment/email-drafts/{retry['_id']}/approve-and-send",
        headers=admin_headers,
    )
    assert sent.status_code == 200
    assert sent.json()["delivery_status"] == "sent"


def test_recruitment_assistant_requires_hr_authentication(client):
    response = client.post("/api/v1/recruitment/assistant", json={"message": "Show applicants"})
    assert response.status_code == 401


def test_recruitment_assistant_reports_grounded_empty_state(client, admin_headers):
    response = client.post(
        "/api/v1/recruitment/assistant",
        headers=admin_headers,
        json={"message": "Show applicants"},
    )
    assert response.status_code == 503
    assert "LLM_API_KEY" in response.json()["detail"]


class FakeLLM:
    """Scripted provider-neutral client: each step is a callable (messages, tools) -> LLMResponse."""

    def __init__(self, *steps):
        self.steps = list(steps)
        self.calls = []

    def chat(self, messages, tools=None):
        self.calls.append(messages)
        return self.steps.pop(0)(messages, tools)


def tool_step(name, arguments):
    return lambda _m, _t: LLMResponse(content="", tool_calls=[ToolCall(id=f"call-{name}", name=name, arguments=arguments)])


def text_step(text):
    return lambda _m, _t: LLMResponse(content=text)


def test_recruitment_assistant_uses_authorized_database_tool(client, admin_headers, monkeypatch):
    application_id = create_application(client, admin_headers, monkeypatch)
    fake = FakeLLM(
        tool_step("get_application_match_details", f'{{"application_id":"{application_id}"}}'),
        text_step("The application is in needs review."),
    )
    monkeypatch.setattr(recruitment_agent, "get_llm_client", lambda: fake)
    response = client.post(
        "/api/v1/recruitment/assistant",
        headers=admin_headers,
        json={"message": "Show this candidate's matching details", "application_id": application_id},
    )
    assert response.status_code == 200, response.text
    payload = response.json()
    assert payload["source"] == "configured_ai_agent"
    assert "needs review" in payload["message"]
    assert payload["tool_calls"][0]["name"] == "get_application_match_details"
    assert payload["tool_calls"][0]["result"]["application_id"] == application_id


def test_agent_never_reports_booking_when_no_schedule_plan_tool_ran(client, admin_headers, monkeypatch):
    monkeypatch.setattr(recruitment_agent, "_execute_tool", lambda *_args, **_kwargs: {"count": 0, "results": []})
    fake = FakeLLM(tool_step("search_applications", "{}"), text_step("All interviews were booked successfully."))
    monkeypatch.setattr(recruitment_agent, "get_llm_client", lambda: fake)
    response = client.post(
        "/api/v1/recruitment/assistant",
        headers=admin_headers,
        json={"message": "Schedule interviews for the shortlisted applicants"},
    )
    assert response.status_code == 200, response.text
    assert "No schedule plan was created" in response.json()["message"]
    assert "booked successfully" not in response.json()["message"]


def test_agent_refuses_non_hr_requests_without_calling_the_model(client, admin_headers, monkeypatch):
    def forbidden():
        raise AssertionError("The LLM must not be called for off-topic requests.")

    monkeypatch.setattr(recruitment_agent, "get_llm_client", forbidden)
    for message in ("Write me a poem about the sea", "What is the capital of France?", "Give me a pasta recipe please"):
        reply = client.post("/api/v1/recruitment/assistant", headers=admin_headers, json={"message": message})
        assert reply.status_code == 200, reply.text
        assert reply.json()["source"] == "scope_guard"
        assert "HR tasks" in reply.json()["message"]
        assert reply.json()["tool_calls"] == []
    greeting = client.post("/api/v1/recruitment/assistant", headers=admin_headers, json={"message": "hello"}).json()
    assert greeting["source"] == "scope_guard"


def test_agent_hr_tools_create_summarize_and_guard_status_changes(client, admin_headers, monkeypatch):
    application_id = create_application(client, admin_headers, monkeypatch)
    monkeypatch.setattr(
        recruitment_agent,
        "get_llm_client",
        lambda: FakeLLM(
            tool_step("create_job_draft", '{"title":"Support Analyst","required_skills":["Excel"]}'),
            text_step("Created."),
        ),
    )
    created = client.post(
        "/api/v1/recruitment/assistant", headers=admin_headers, json={"message": "Please create a job posting for a Support Analyst"}
    ).json()
    assert created["tool_calls"][0]["result"]["status"] == "draft"
    assert any(job["title"] == "Support Analyst" for job in client.get("/api/v1/recruitment/jobs", headers=admin_headers).json())

    user = {"_id": "admin-id"}
    from app.services import agent_tools

    from app.core.database import get_database

    database = next(client.app.dependency_overrides[get_database]())
    summary = agent_tools.hiring_summary(database, {}, None)
    assert summary["applications_by_stage"].get("needs_review", 0) >= 1
    refused = agent_tools.set_application_stage(database, user, {"application_id": application_id, "stage": "interview"}, "move them to interview")
    assert "schedule" in refused["error"].lower()
    no_verb = agent_tools.set_application_stage(database, user, {"application_id": application_id, "stage": "screening"}, "tell me about them")
    assert "explicit" in no_verb["error"]
    moved = agent_tools.set_application_stage(database, user, {"application_id": application_id, "stage": "screening"}, "move this candidate to screening")
    assert moved["stage"] == "screening"


def test_email_draft_can_be_edited_deleted_and_recreated(client, admin_headers, monkeypatch):
    application_id = create_application(client, admin_headers, monkeypatch)
    draft = client.post(
        f"/api/v1/recruitment/applications/{application_id}/email-drafts",
        headers=admin_headers,
        json={"email_type": "shortlist", "message": "Original note."},
    ).json()
    edited = client.patch(
        f"/api/v1/recruitment/email-drafts/{draft['_id']}",
        headers=admin_headers,
        json={"subject": "Updated subject", "body": "Updated body."},
    )
    assert edited.status_code == 200, edited.text
    assert edited.json()["subject"] == "Updated subject"
    assert edited.json()["body"] == "Updated body."
    assert client.get(
        f"/api/v1/recruitment/email-drafts/{draft['_id']}/status",
        headers=admin_headers,
    ).json()["delivery_status"] == "not_sent"

    deleted = client.delete(f"/api/v1/recruitment/email-drafts/{draft['_id']}", headers=admin_headers)
    assert deleted.status_code == 204
    assert client.get(
        f"/api/v1/recruitment/email-drafts/{draft['_id']}",
        headers=admin_headers,
    ).status_code == 404
    assert client.get(
        f"/api/v1/recruitment/applications/{application_id}/email-drafts",
        headers=admin_headers,
    ).json() == []
    recreated = client.post(
        f"/api/v1/recruitment/applications/{application_id}/email-drafts",
        headers=admin_headers,
        json={"email_type": "shortlist", "message": "New draft."},
    )
    assert recreated.status_code == 201, recreated.text


def test_interview_invitation_requires_a_persisted_scheduled_interview(client, admin_headers, monkeypatch):
    application_id = create_application(client, admin_headers, monkeypatch)
    response = client.post(
        f"/api/v1/recruitment/applications/{application_id}/email-drafts",
        headers=admin_headers,
        json={"email_type": "interview_invitation", "message": ""},
    )
    assert response.status_code == 409
    assert "Schedule an interview" in response.json()["detail"]
