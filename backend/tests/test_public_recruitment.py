from datetime import datetime, timedelta, timezone
from io import BytesIO

from docx import Document


def make_docx() -> bytes:
    document = Document()
    document.add_paragraph("Taylor Candidate")
    document.add_paragraph("taylor@example.com")
    document.add_paragraph("Experience")
    document.add_paragraph("Built dependable APIs with Python and PostgreSQL for five years.")
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def published_job(client, headers, **fields):
    payload = {
        "title": "Backend Developer",
        "description": "Build distributed backend services.",
        "responsibilities": ["Design APIs", "Improve service reliability"],
        "required_skills": ["Python"],
        "mandatory_criteria": ["work authorization"],
        **fields,
    }
    response = client.post("/api/v1/recruitment/jobs", headers=headers, json=payload)
    assert response.status_code == 201, response.text
    return response.json(), client.post(f"/api/v1/recruitment/jobs/{response.json()['id']}/publish", headers=headers)


def application_form(token, **overrides):
    data = {
        "name": "Taylor Candidate",
        "email": "taylor@example.com",
        "phone": "555-123-4567",
        "education": "Bachelor degree in Computer Science",
        "experience_years": "5",
        "answers": "{}",
    }
    data.update(overrides)
    return data


def test_public_job_publish_form_and_duplicate_application(client, admin_headers):
    draft, publish = published_job(
        client,
        admin_headers,
        questions=[{"key": "why_role", "label": "Why this role?", "required": True}],
    )
    assert publish.status_code == 200
    token = publish.json()["public_url"].rsplit("/", 1)[1]
    public_job = client.get(f"/api/v1/recruitment/public/jobs/{token}")
    assert public_job.status_code == 200
    assert public_job.json()["title"] == "Backend Developer"
    assert "id" not in public_job.json()
    missing_answer = client.post(
        f"/api/v1/recruitment/public/jobs/{token}/applications",
        data=application_form(token),
        files={"resume": ("resume.docx", make_docx(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert missing_answer.status_code == 422
    submitted = client.post(
        f"/api/v1/recruitment/public/jobs/{token}/applications",
        data=application_form(token, answers='{"why_role":"I enjoy service design."}'),
        files={"resume": ("resume.docx", make_docx(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert submitted.status_code == 202, submitted.text
    assert submitted.json()["status"] == "submitted"
    duplicate = client.post(
        f"/api/v1/recruitment/public/jobs/{token}/applications",
        data=application_form(token, answers='{"why_role":"Another attempt."}'),
        files={"resume": ("resume.docx", make_docx(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert duplicate.status_code == 409
    app_list = client.get("/api/v1/recruitment/applications", headers=admin_headers)
    assert len(app_list.json()) == 1
    assert app_list.json()[0]["matching_status"] == "pending"


def test_public_application_rejects_draft_closed_and_expired_jobs(client, admin_headers):
    response = client.post(
        "/api/v1/recruitment/jobs",
        headers=admin_headers,
        json={"title": "Draft Engineer", "description": "A draft role"},
    )
    token = response.json()["id"]
    assert client.get(f"/api/v1/recruitment/public/jobs/{token}").status_code == 404
    job, published = published_job(
        client,
        admin_headers,
        title="Expired Engineer",
        deadline=(datetime.now(timezone.utc) - timedelta(days=1)).isoformat(),
    )
    token = published.json()["public_url"].rsplit("/", 1)[1]
    expired = client.post(
        f"/api/v1/recruitment/public/jobs/{token}/applications",
        data=application_form(token),
        files={"resume": ("resume.docx", make_docx(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert expired.status_code == 410
    closed = client.post(f"/api/v1/recruitment/jobs/{job['id']}/close", headers=admin_headers)
    assert closed.status_code == 200
    assert client.get(f"/api/v1/recruitment/public/jobs/{token}").status_code == 404


def test_job_scoring_weights_must_total_one_hundred(client, admin_headers):
    response = client.post(
        "/api/v1/recruitment/jobs",
        headers=admin_headers,
        json={"title": "Bad Weights", "scoring_weights": {"skills": 10}},
    )
    assert response.status_code == 422
