import base64
import json
import re
from datetime import datetime, timedelta, timezone
from io import BytesIO

import httpx
import pytest
from docx import Document

from app.api.v1 import interview_workflow
from app.core.config import settings
from app.core.database import get_database
from app.services import interview_notifications, meetings, shortlisting
from app.services.llm import (
    LLMConfigError,
    LLMNotConfigured,
    LLMResponse,
    get_llm_client,
    validate_llm_config,
)
from app.services.llm.anthropic_provider import AnthropicClient, to_anthropic_messages
from app.services.llm.openai_provider import OpenAICompatibleClient
from test_interview_workflow import AvailableCalendar, configure_rounds, create_application, create_job, record_round_one


def database(client):
    return next(client.app.dependency_overrides[get_database]())


def post_resume(client, job_title, email, name, text, filename="resume.txt", content=None, content_type="text/plain"):
    payload = content if content is not None else f"{name}\n{email}\n{text}".encode()
    response = client.post(
        "/api/v1/recruitment/google-forms/applications",
        headers={"X-Integration-Key": "test-integration-secret"},
        json={
            "name": name, "email": email, "job_title": job_title, "resume_filename": filename,
            "resume_content_type": content_type, "resume_base64": base64.b64encode(payload).decode(),
        },
    )
    assert response.status_code == 201, response.text
    return response.json()["id"]


# ---------------------------------------------------------------- LLM provider abstraction


def test_llm_provider_is_selected_by_configuration_only(monkeypatch):
    for name in ("llm_provider", "llm_model", "llm_api_key", "llm_base_url", "agent_api_key", "agent_api_base_url"):
        monkeypatch.setattr(settings, name, None)
    with pytest.raises(LLMNotConfigured):
        get_llm_client()

    monkeypatch.setattr(settings, "llm_provider", "anthropic")
    monkeypatch.setattr(settings, "llm_model", "claude-test")
    monkeypatch.setattr(settings, "llm_api_key", "k")
    assert isinstance(get_llm_client(), AnthropicClient)

    monkeypatch.setattr(settings, "llm_provider", "openai")
    assert isinstance(get_llm_client(), OpenAICompatibleClient)

    monkeypatch.setattr(settings, "llm_provider", "openai_compatible")
    monkeypatch.setattr(settings, "llm_api_key", None)
    monkeypatch.setattr(settings, "llm_base_url", "http://localhost:11434/v1")
    assert isinstance(get_llm_client(), OpenAICompatibleClient)  # local server: no key needed


def test_llm_configuration_errors_are_clear(monkeypatch):
    monkeypatch.setattr(settings, "llm_base_url", None)
    monkeypatch.setattr(settings, "agent_api_base_url", None)
    monkeypatch.setattr(settings, "llm_provider", "mystery")
    monkeypatch.setattr(settings, "llm_model", "m")
    monkeypatch.setattr(settings, "llm_api_key", "k")
    assert "not supported" in validate_llm_config()[0]
    monkeypatch.setattr(settings, "llm_provider", "openai")
    monkeypatch.setattr(settings, "llm_model", None)
    monkeypatch.setattr(settings, "agent_model", "")
    monkeypatch.setattr(settings, "llm_api_key", None)
    monkeypatch.setattr(settings, "agent_api_key", None)
    monkeypatch.setattr(settings, "llm_provider", "openai")
    problems = validate_llm_config()
    assert any("LLM_MODEL" in item for item in problems) and any("LLM_API_KEY" in item for item in problems)
    with pytest.raises(LLMConfigError):
        get_llm_client()


def test_legacy_agent_settings_still_configure_openai(monkeypatch):
    for name in ("llm_provider", "llm_model", "llm_api_key", "llm_base_url"):
        monkeypatch.setattr(settings, name, None)
    monkeypatch.setattr(settings, "agent_api_key", "legacy")
    monkeypatch.setattr(settings, "agent_api_base_url", None)
    monkeypatch.setattr(settings, "agent_model", "gpt-test")
    assert validate_llm_config() == []
    assert isinstance(get_llm_client(), OpenAICompatibleClient)


def test_anthropic_adapter_translates_messages_and_tool_calls(monkeypatch):
    system, converted = to_anthropic_messages([
        {"role": "system", "content": "rules"},
        {"role": "user", "content": "hi"},
        {"role": "assistant", "content": "", "tool_calls": [{"id": "t1", "type": "function", "function": {"name": "list_jobs", "arguments": "{}"}}]},
        {"role": "tool", "tool_call_id": "t1", "content": "{\"count\": 0}"},
    ])
    assert system == "rules"
    assert converted[1]["content"][0]["type"] == "tool_use"
    assert converted[2] == {"role": "user", "content": [{"type": "tool_result", "tool_use_id": "t1", "content": "{\"count\": 0}"}]}

    def fake_post(url, json, headers, timeout):
        assert url.endswith("/v1/messages") and headers["x-api-key"] == "k"
        assert json["tools"][0]["name"] == "list_jobs"
        return httpx.Response(200, json={"content": [{"type": "text", "text": "ok"}, {"type": "tool_use", "id": "u1", "name": "list_jobs", "input": {"status": "draft"}}]}, request=httpx.Request("POST", url))

    monkeypatch.setattr(httpx, "post", fake_post)
    client = AnthropicClient(api_key="k", model="m", base_url=None, temperature=0, max_tokens=100, timeout=5)
    reply = client.chat([{"role": "user", "content": "x"}], [{"type": "function", "function": {"name": "list_jobs", "description": "d", "parameters": {"type": "object", "properties": {}}}}])
    assert reply.content == "ok"
    assert reply.tool_calls[0].name == "list_jobs" and json.loads(reply.tool_calls[0].arguments) == {"status": "draft"}


# ---------------------------------------------------------------- meetings


def test_meeting_link_is_never_invented_and_provider_is_configurable(client, admin_headers, monkeypatch):
    db = database(client)
    job_id = create_job(client, admin_headers)
    configure_rounds(client, admin_headers, job_id)
    application = create_application(client, job_id, "m@example.com", "Meet Person")
    record_round_one(client, admin_headers, application, "passed")
    interview = db["interviews"].find_one({"application_id": application})
    candidate = {"name": "Meet Person"}
    job = {"title": "Workflow Engineer"}
    interview["starts_at"] = datetime.now(timezone.utc) + timedelta(days=1)

    meetings.ensure_meeting(db, interview, candidate, job)
    assert interview["meeting_status"] == "not_configured" and interview["meeting_url"] is None

    calls = []
    monkeypatch.setattr(settings, "meeting_provider", "webhook")
    monkeypatch.setattr(settings, "meeting_webhook_url", "https://meetings.example/hook")
    monkeypatch.setattr(
        httpx, "post",
        lambda url, **kw: calls.append(kw["json"]["action"]) or httpx.Response(200, json={"join_url": "https://meet.example/abc", "meeting_id": "m-1"}, request=httpx.Request("POST", url)),
    )
    meetings.ensure_meeting(db, interview, candidate, job)
    assert interview["meeting_url"] == "https://meet.example/abc" and interview["meeting_status"] == "created"
    assert interview["meeting_expires_at"] > interview["starts_at"]
    meetings.ensure_meeting(db, interview, candidate, job)
    assert calls == ["create"]  # idempotent

    meetings.invalidate_meeting(db, interview, "cancelled")
    stored = db["interviews"].find_one({"_id": interview["_id"]})
    assert stored["meeting_status"] == "invalidated" and stored["meeting_url"] is None and calls[-1] == "delete"


def test_unsafe_meeting_links_and_bad_configuration_are_rejected(monkeypatch):
    monkeypatch.setattr(settings, "meeting_provider", "zoom")
    monkeypatch.setattr(settings, "zoom_account_id", None)
    assert "ZOOM_ACCOUNT_ID" in meetings.validate_meeting_config()[0]
    monkeypatch.setattr(settings, "meeting_provider", "carrier-pigeon")
    assert "not supported" in meetings.validate_meeting_config()[0]
    monkeypatch.setattr(settings, "meeting_provider", "webhook")
    monkeypatch.setattr(settings, "meeting_webhook_url", "https://meetings.example/hook")
    monkeypatch.setattr(
        httpx, "post",
        lambda url, **kw: httpx.Response(200, json={"join_url": "http://insecure.example/x"}, request=httpx.Request("POST", url)),
    )
    with pytest.raises(meetings.MeetingProviderError):
        meetings.create_meeting(topic="t", start=datetime.now(timezone.utc), duration_minutes=30, timezone_name="UTC", interview_id="i")


# ---------------------------------------------------------------- interviewer evaluation links and next-round eligibility


def three_round_config(client, headers, job_id):
    rounds = [
        {"round_id": "technical-1", "name": "Technical Interview 1", "order": 1, "required_interviewers": ["interviewer@example.com"]},
        {"round_id": "technical-2", "name": "Technical Interview 2", "order": 2, "required_interviewers": ["interviewer@example.com"], "prerequisite_round_ids": ["technical-1"]},
        {"round_id": "technical-3", "name": "Technical Interview 3", "order": 3, "required_interviewers": ["interviewer@example.com"], "prerequisite_round_ids": ["technical-2"]},
    ]
    assert client.put(f"/api/v1/interviews/jobs/{job_id}/rounds", headers=headers, json={"rounds": rounds}).status_code == 200


def schedule_round(client, headers, job_id, instruction):
    plan = client.post("/api/v1/interviews/assistant", headers=headers, json={"job_id": job_id, "instruction": instruction}).json()
    assert plan["counts"]["awaiting_confirmation"] >= 1, plan
    confirmed = client.post(f"/api/v1/interviews/scheduling-plans/{plan['plan_id']}/confirm", headers=headers)
    assert confirmed.status_code == 200, confirmed.text
    return plan, confirmed.json()


def test_evaluation_links_drive_outcomes_and_next_round_eligibility(client, admin_headers, monkeypatch):
    mails = []
    monkeypatch.setattr(interview_notifications, "send_email", lambda **mail: mails.append(mail))
    monkeypatch.setattr(interview_workflow, "GoogleCalendar", lambda: AvailableCalendar())
    db = database(client)
    job_id = create_job(client, admin_headers)
    three_round_config(client, admin_headers, job_id)
    people = {
        "passed": create_application(client, job_id, "yes@example.com", "Yes Person"),
        "failed": create_application(client, job_id, "no@example.com", "Nope Person"),
        "on_hold": create_application(client, job_id, "wait@example.com", "Wait Person"),
    }
    for application in people.values():
        record_round_one(client, admin_headers, application, "passed")

    _plan, confirmed = schedule_round(client, admin_headers, job_id, "Schedule Technical Interview 2 for all")
    assert confirmed["counts"]["scheduled"] == 3
    # Emailing and scheduling never mark anyone as passed.
    assert {i["outcome"] for i in db["interviews"].find({"round_id": "technical-2"})} == {"pending_evaluation"}

    interviewer_mails = [m for m in mails if m["recipient"] == "interviewer@example.com"]
    assert len(interviewer_mails) == 3
    tokens = {}
    for mail in interviewer_mails:
        body = mail["body"]
        link = re.search(r"/evaluate/([\w-]+)", body).group(1)
        name = re.search(r"Name: (.+)", body).group(1)
        tokens[name] = link
        assert "Workflow Engineer" in body and "Build dependable software." in body  # job title and description
        assert "yes@example.com" in body or "no@example.com" in body or "wait@example.com" in body
        assert "UTC" in body and "/evaluations/" in body  # timezone and secure resume link
        assert mail["attachments"] and mail["attachments"][0][0].endswith(".txt")  # resume attached
    candidate_mail = next(m for m in mails if m["recipient"] == "yes@example.com")
    assert "passed Technical Interview 1" in candidate_mail["body"]
    assert "could not be generated" in candidate_mail["body"]  # no provider configured: nothing invented
    assert "http" not in candidate_mail["body"].split("Meeting link")[1].split("\n")[0]

    token = tokens["Yes Person"]
    page = client.get(f"/api/v1/evaluations/{token}")
    assert page.status_code == 200
    assert page.json()["candidate"]["name"] == "Yes Person" and page.json()["round"]["name"] == "Technical Interview 2"
    assert client.get("/api/v1/evaluations/not-a-real-token").status_code == 404
    assert client.get(f"/api/v1/evaluations/{token}/resume").status_code == 200

    early = client.post(f"/api/v1/evaluations/{token}", json={"outcome": "passed", "feedback": "Strong problem solving."})
    assert early.status_code == 409 and "opens when the interview starts" in early.json()["detail"]
    db["interviews"].update_many({"round_id": "technical-2"}, {"$set": {"starts_at": datetime.now(timezone.utc) - timedelta(hours=1)}})

    assert client.post(f"/api/v1/evaluations/{token}", json={"outcome": "passed", "feedback": "short"}).status_code == 422
    assert client.post(f"/api/v1/evaluations/{token}", json={"outcome": "maybe", "feedback": "Long enough feedback."}).status_code == 422
    first = client.post(f"/api/v1/evaluations/{token}", json={"outcome": "passed", "feedback": "Strong problem solving and design."})
    assert first.status_code == 200, first.text
    duplicate = client.post(f"/api/v1/evaluations/{token}", json={"outcome": "failed", "feedback": "Changed my mind later."})
    assert duplicate.status_code == 409 and "already recorded" in duplicate.json()["detail"]

    assert client.post(f"/api/v1/evaluations/{tokens['Nope Person']}", json={"outcome": "failed", "feedback": "Could not explain basics."}).status_code == 200
    held = client.post(f"/api/v1/evaluations/{tokens['Wait Person']}", json={"outcome": "on_hold", "feedback": "Need a second opinion first."})
    assert held.status_code == 200 and "pending" in held.json()["message"]

    applications = {a["id"]: a for a in client.get("/api/v1/recruitment/applications", headers=admin_headers).json()}
    assert "Passed Technical Interview 2 \u00b7 eligible for Technical Interview 3" in applications[people["passed"]]["pipeline_status"]
    assert applications[people["failed"]]["pipeline_status"] == "Not selected after Technical Interview 2"
    assert applications[people["on_hold"]]["pipeline_status"] == "On hold after Technical Interview 2"

    plan = client.post("/api/v1/interviews/assistant", headers=admin_headers, json={"job_id": job_id, "instruction": "Schedule Technical Interview 3 for everyone"}).json()
    status = {item["candidate_email"]: item for item in plan["items"]}
    assert status["yes@example.com"]["status"] == "awaiting_confirmation"
    assert status["no@example.com"]["status"] == "ineligible" and "failed" in status["no@example.com"]["reason"]
    assert status["wait@example.com"]["status"] == "ineligible" and "on hold" in status["wait@example.com"]["reason"]

    # An on-hold interview can be resolved from the same link; a final outcome cannot be changed.
    resolved = client.post(f"/api/v1/evaluations/{tokens['Wait Person']}", json={"outcome": "passed", "feedback": "Second opinion was positive."})
    assert resolved.status_code == 200
    assert client.post(f"/api/v1/evaluations/{tokens['Wait Person']}", json={"outcome": "failed", "feedback": "Trying to flip it again."}).status_code == 409


def test_evaluation_links_expire_and_die_on_cancellation(client, admin_headers, monkeypatch):
    mails = []
    monkeypatch.setattr(interview_notifications, "send_email", lambda **mail: mails.append(mail))
    monkeypatch.setattr(interview_workflow, "GoogleCalendar", lambda: AvailableCalendar())
    db = database(client)
    job_id = create_job(client, admin_headers)
    configure_rounds(client, admin_headers, job_id)
    application = create_application(client, job_id, "gone@example.com", "Gone Person")
    record_round_one(client, admin_headers, application, "passed")
    schedule_round(client, admin_headers, job_id, "Schedule Technical Interview 2")
    token = re.search(r"/evaluate/([\w-]+)", next(m for m in mails if m["recipient"] == "interviewer@example.com")["body"]).group(1)
    assert client.get(f"/api/v1/evaluations/{token}").status_code == 200

    db["evaluation_tokens"].update_many({}, {"$set": {"expires_at": datetime.now(timezone.utc) - timedelta(minutes=1)}})
    assert client.get(f"/api/v1/evaluations/{token}").status_code == 410
    db["evaluation_tokens"].update_many({}, {"$set": {"expires_at": datetime.now(timezone.utc) + timedelta(days=1)}})

    interview = next(i for i in client.get("/api/v1/interviews", headers=admin_headers).json() if i["round_id"] == "technical-2")
    payload = {k: v for k, v in interview.items() if k not in {"id", "candidate_name", "candidate_email", "job_title", "pipeline_status", "email_status", "meeting_status", "meeting_url", "meeting_error"}}
    payload["status"] = "cancelled"
    assert client.put(f"/api/v1/interviews/{interview['id']}", headers=admin_headers, json=payload).status_code == 200
    gone = client.get(f"/api/v1/evaluations/{token}")
    assert gone.status_code == 410 and "rescheduled or cancelled" in gone.json()["detail"]
    assert client.post(f"/api/v1/evaluations/{token}", json={"outcome": "passed", "feedback": "Should not be accepted."}).status_code == 410


def test_round_cannot_be_scheduled_from_instruction_wording_alone(client, admin_headers, monkeypatch):
    monkeypatch.setattr(interview_workflow, "GoogleCalendar", lambda: AvailableCalendar())
    job_id = create_job(client, admin_headers)
    configure_rounds(client, admin_headers, job_id)
    create_application(client, job_id, "nopass@example.com", "No Pass")
    plan = client.post(
        "/api/v1/interviews/assistant", headers=admin_headers,
        json={"job_id": job_id, "instruction": "Schedule Technical Interview 2 for those who passed Technical Interview 1"},
    ).json()
    assert plan["counts"]["eligible"] == 0
    assert plan["items"][0]["status"] == "ineligible"


def test_scheduling_works_without_google_calendar_and_creates_no_events(client, admin_headers, monkeypatch):
    monkeypatch.setattr(interview_notifications, "send_email", lambda **mail: None)
    job_id = create_job(client, admin_headers)
    configure_rounds(client, admin_headers, job_id)
    application = create_application(client, job_id, "plain@example.com", "Plain Person")
    record_round_one(client, admin_headers, application, "passed")
    plan, confirmed = schedule_round(client, admin_headers, job_id, "Schedule Technical Interview 2")
    assert plan["availability_sources"] == {"hrms_bookings": True, "google_calendar_read_only": False}
    assert "HRMS bookings only" in plan["availability_notice"]
    stored = database(client)["interviews"].find_one({"round_id": "technical-2"})
    assert "calendar_event_id" not in stored and stored["status"] == "scheduled"
    # HRMS bookings still prevent double-booking the interviewer.
    second = create_application(client, job_id, "plain2@example.com", "Plain Two")
    record_round_one(client, admin_headers, second, "passed")
    second_plan = client.post("/api/v1/interviews/assistant", headers=admin_headers, json={"job_id": job_id, "instruction": "Schedule Technical Interview 2"}).json()
    start_one = datetime.fromisoformat(plan["items"][0]["starts_at"])
    start_two = datetime.fromisoformat(next(i for i in second_plan["items"] if i["candidate_email"] == "plain2@example.com")["starts_at"])
    assert start_two != start_one


# ---------------------------------------------------------------- resume parsing and shortlisting


def docx_with_table() -> bytes:
    document = Document()
    document.add_paragraph("Asha Verma")
    document.add_paragraph("TECHNICAL SKILLS")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Languages"
    table.rows[0].cells[1].text = "JS, TS, Py"
    document.add_paragraph("Academic Projects")
    document.add_paragraph("Built an ML powered chatbot using LangChain and RESTful APIs")
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def test_parser_reads_docx_tables_and_flexible_headings():
    from app.services.resume_parser import extract_text, parse_resume

    text = extract_text("r.docx", docx_with_table())
    assert "JS, TS, Py" in text
    parsed = parse_resume(text)
    assert "javascript" in parsed["skills"] and "typescript" in parsed["skills"] and "python" in parsed["skills"]
    assert "LangChain" in parsed["projects"]


def test_unreadable_resumes_report_a_reason():
    from app.services.resume_parser import extract_text

    for name, data in (("a.docx", b"not a zip"), ("a.pdf", b"%PDF-broken"), ("a.txt", b"   \n ")):
        with pytest.raises(ValueError):
            extract_text(name, data)


JOB = {
    "title": "AI Automation Intern", "description": "Automate hiring workflows with language models.",
    "required_skills": ["JavaScript", "Machine Learning", "REST API"], "preferred_skills": ["Docker", "LangChain"],
    "responsibilities": ["Build automation workflows using language models", "Integrate REST APIs"],
    "mandatory_criteria": [],
}


def test_assessment_uses_synonyms_and_separates_missing_evidence_from_lacking():
    synonyms = shortlisting.assess_candidate(
        JOB, {"experience_years": None, "projects": "Built an ML chatbot with LangChain exposed via RESTful APIs in JS", "internships": ""},
        "SKILLS\nJS, ML, RESTful APIs\nPROJECTS\nBuilt an ML chatbot with LangChain exposed via RESTful APIs in JS",
    )
    statuses = {item["requirement"]: item for item in synonyms["required_skills"]}
    assert all(item["status"] in ("met", "related") for item in statuses.values())
    assert statuses["JavaScript"]["method"] == "synonym" and statuses["JavaScript"]["evidence"]
    assert synonyms["eligibility"] == "eligible"
    assert any(item["requirement"] == "LangChain" and item["status"] == "met" for item in synonyms["preferred_skills"])

    unrelated = shortlisting.assess_candidate(JOB, {"projects": "", "internships": ""}, "SKILLS\nBookkeeping, Tally\nEXPERIENCE\nAccounts assistant")
    assert unrelated["eligibility"] == "needs_review"  # never rejected on wording alone
    assert all(item["status"] == "not_evidenced" for item in unrelated["required_skills"])
    assert "absence of evidence" in unrelated["explanation"]
    assert all(gap["kind"] == "missing_evidence" for gap in unrelated["evidence_gaps"])


def test_llm_evaluator_claims_must_be_quoted_from_the_resume():
    class Llm:
        def chat(self, messages, tools=None):
            return LLMResponse(content=json.dumps({"requirements": [
                {"name": "Machine Learning", "status": "met", "quote": "trained a neural network classifier"},
                {"name": "REST API", "status": "met", "quote": "led a team of fifty engineers"},
                {"name": "JavaScript", "status": "contradicted", "quote": "no experience with JavaScript"},
            ], "summary": "Relevant ML work."}))

    result = shortlisting.assess_candidate(
        JOB, {"projects": "", "internships": ""},
        "PROJECTS\nI trained a neural network classifier for images.\nNote: no experience with JavaScript yet.", llm=Llm(),
    )
    by_name = {item["requirement"]: item for item in result["required_skills"]}
    assert by_name["Machine Learning"]["status"] == "related" and by_name["Machine Learning"]["method"] == "llm"
    assert by_name["REST API"]["status"] == "not_evidenced"  # invented quote ignored
    assert by_name["JavaScript"]["status"] == "contradicted"  # only because the resume says so
    assert result["llm_used"] is True and result["eligibility"] in ("needs_review", "not_eligible")


def test_semantic_embeddings_can_settle_requirements_but_remain_configurable(monkeypatch):
    class Embedder:
        def embed(self, texts):
            return [[1.0, 0.0] if ("neural" in t.casefold() or "machine learning" in t.casefold()) else [0.0, 1.0] for t in texts]

    result = shortlisting.assess_candidate(JOB, {"projects": "", "internships": ""}, "PROJECTS\nTrained neural networks for vision tasks.", embedder=Embedder())
    ml = next(i for i in result["required_skills"] if i["requirement"] == "Machine Learning")
    assert ml["status"] == "related" and ml["method"] == "semantic" and result["semantic_used"]
    monkeypatch.setattr(settings, "shortlist_semantic_threshold", 1.0)

    class Orthogonal:
        def embed(self, texts):
            return [[1.0, 0.0] if "neural" in text.casefold() else [0.0, 1.0] for text in texts]

    strict = shortlisting.assess_candidate(JOB, {"projects": "", "internships": ""}, "PROJECTS\nTrained neural networks for vision tasks.", embedder=Orthogonal())
    assert next(i for i in strict["required_skills"] if i["requirement"] == "Machine Learning")["status"] == "not_evidenced"


def test_weights_and_thresholds_are_configurable(monkeypatch):
    base = shortlisting.assess_candidate(JOB, {"projects": "", "internships": ""}, "SKILLS\nJS, ML\nPROJECTS\nBuilt an ML bot in JS")
    monkeypatch.setattr(settings, "shortlist_min_score", 99.0)
    strict = shortlisting.assess_candidate(JOB, {"projects": "", "internships": ""}, "SKILLS\nJS, ML\nPROJECTS\nBuilt an ML bot in JS")
    assert base["eligibility"] != strict["eligibility"] or strict["eligibility"] != "eligible"
    assert strict["thresholds"]["min_score"] == 99.0
    only_required = shortlisting.assess_candidate(JOB, {"projects": "", "internships": ""}, "SKILLS\nJS, ML, REST\n", weights={"required_skills": 100.0})
    assert only_required["weights_applied"] == {"required_skills": 100.0}
    with pytest.raises(ValueError):
        type(settings)(shortlist_weights={"required_skills": 50})


def test_shortlist_endpoint_ranks_limits_and_reports_unreadable_resumes(client, admin_headers):
    created = client.post(
        "/api/v1/recruitment/jobs", headers=admin_headers,
        json={**{k: v for k, v in JOB.items() if k != "mandatory_criteria"}},
    )
    assert created.status_code == 201, created.text
    job_id = created.json()["id"]
    client.post(f"/api/v1/recruitment/jobs/{job_id}/publish", headers=admin_headers)
    strong = post_resume(client, JOB["title"], "strong@example.com", "Strong One", "SKILLS\nJS, ML, RESTful APIs, Docker\nPROJECTS\nBuilt LangChain automation workflows with ML models and REST APIs")
    ok = post_resume(client, JOB["title"], "ok@example.com", "Okay Two", "SKILLS\nJavaScript, Machine Learning, REST API\nPROJECTS\nA small ML demo")
    other = post_resume(client, JOB["title"], "other@example.com", "Other Three", "SKILLS\nTally, Bookkeeping\nEXPERIENCE\nAccounts assistant")
    broken = post_resume(client, JOB["title"], "broken@example.com", "Broken Four", "SKILLS\nJS")
    from pathlib import Path

    Path(database(client)["applications"].find_one({"_id": broken})["resume_path"]).write_bytes(b"\xff\xfe\x00 damaged")

    result = client.post(f"/api/v1/recruitment/jobs/{job_id}/shortlist", headers=admin_headers, json={"limit": 1, "apply": True})
    assert result.status_code == 200, result.text
    body = result.json()
    by_id = {row["application_id"]: row for row in body["results"]}
    assert by_id[strong]["status"] == "shortlisted" and by_id[strong]["rank"] == 1
    assert by_id[ok]["status"] == "eligible" and by_id[ok]["rank"] == 2  # eligible, beyond the configured limit
    assert by_id[other]["status"] == "needs_review"
    assert by_id[broken]["status"] == "parse_failed" and "unreadable" in by_id[broken]["explanation"]
    assert body["moved_to_screening"] == 1

    detail = client.get(f"/api/v1/recruitment/applications/{strong}", headers=admin_headers).json()
    assert detail["stage"] == "screening" and detail["shortlist_status"] == "shortlisted"
    assert detail["assessment"]["required_skills"][0]["evidence"]
    assert client.get(f"/api/v1/recruitment/applications/{other}", headers=admin_headers).json()["stage"] == "needs_review"
    assert client.get(f"/api/v1/recruitment/jobs/{job_id}/shortlist", headers=admin_headers).json()["results"][0]["application_id"] == strong


# ---------------------------------------------------------------- chatbot context and ambiguity


def test_chat_resolves_these_candidates_from_conversation_context(client, admin_headers, monkeypatch):
    monkeypatch.setattr(interview_notifications, "send_email", lambda **mail: None)
    monkeypatch.setattr(interview_workflow, "GoogleCalendar", lambda: AvailableCalendar())
    job_id = create_job(client, admin_headers)
    configure_rounds(client, admin_headers, job_id)
    for index in range(3):
        application = create_application(client, job_id, f"c{index}@example.com", f"Cand{index} Tester")
        record_round_one(client, admin_headers, application, "passed")
    first = client.post(
        "/api/v1/recruitment/assistant", headers=admin_headers,
        json={"message": "Schedule tech 2 for 2 candidates", "job_id": job_id},
    ).json()
    plan = first["tool_calls"][0]["result"]
    assert len(plan["items"]) == 2
    follow_up = client.post(
        "/api/v1/recruitment/assistant", headers=admin_headers,
        json={
            "message": "Reschedule tech 2 for these candidates",
            "job_id": job_id,
            "context": {"plan_id": plan["plan_id"]},
            "history": [{"role": "user", "content": "Schedule tech 2 for 2 candidates"}, {"role": "assistant", "content": first["message"]}],
        },
    ).json()
    assert {i["candidate_email"] for i in follow_up["tool_calls"][0]["result"]["items"]} == {i["candidate_email"] for i in plan["items"]}


def test_chat_asks_which_job_when_the_request_matches_several(client, admin_headers, monkeypatch):
    for title in ("Backend Developer", "Frontend Developer"):
        assert client.post("/api/v1/recruitment/jobs", headers=admin_headers, json={"title": title, "description": "Build things."}).status_code == 201
    reply = client.post(
        "/api/v1/recruitment/assistant", headers=admin_headers, json={"message": "Schedule technical round 2 for the developer role"}
    ).json()
    result = reply["tool_calls"][0]["result"]
    assert result["status"] == "clarification_required"
    assert {job["title"] for job in result["matching_jobs"]} == {"Backend Developer", "Frontend Developer"}
    assert "Which job" in reply["message"]
